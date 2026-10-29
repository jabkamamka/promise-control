"""Проверка презентации: переполнение текста, выход за границы, пересечения блоков."""
from __future__ import annotations

import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "docs" / "Презентация_7_минут.pptx"

EMU_PER_IN = 914400
CHAR_W = 0.58  # средняя ширина символа в долях кегля (кириллица)


def inches(value: int) -> float:
    return value / EMU_PER_IN


def estimate_height(tf, width_in: float) -> float:
    total = 0.0
    for p in tf.paragraphs:
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
        line_spacing = p.line_spacing or 1.15
        chars = sum(len(r.text) for r in p.runs) or len(p.text)
        per_line = max(1, int(width_in * 72 / (CHAR_W * size)))
        lines = max(1, math.ceil(chars / per_line)) if chars else 1
        total += lines * size * 1.22 * line_spacing / 72
    return total


def main() -> int:
    prs = Presentation(str(PATH))
    slide_w, slide_h = inches(prs.slide_width), inches(prs.slide_height)
    problems = []

    for index, slide in enumerate(prs.slides, start=1):
        boxes = []
        for shape in slide.shapes:
            x, y = inches(shape.left), inches(shape.top)
            w, h = inches(shape.width), inches(shape.height)
            if x < -0.02 or y < -0.02 or x + w > slide_w + 0.02 or y + h > slide_h + 0.02:
                problems.append(f"слайд {index}: фигура вне слайда ({x:.2f},{y:.2f},{w:.2f}x{h:.2f})")

            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            need = estimate_height(shape.text_frame, w)
            if need > h + 0.06:
                problems.append(
                    f"слайд {index}: возможное переполнение «{text[:38]}…» — нужно {need:.2f}\", есть {h:.2f}\""
                )
            boxes.append((x, y, w, h, text[:30]))

    # пересечения текстовых блоков (грубая проверка)
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[i]
            bx, by, bw, bh, bt = boxes[j]
            overlap_x = min(ax + aw, bx + bw) - max(ax, bx)
            overlap_y = min(ay + ah, by + bh) - max(ay, by)
            if overlap_x > 0.12 and overlap_y > 0.12:
                problems.append(f"пересечение блоков: «{at}…» и «{bt}…» ({overlap_x:.2f}x{overlap_y:.2f}\")")

    print(f"слайдов: {len(prs.slides)}, проблем: {len(problems)}")
    for p in problems:
        print(" -", p)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
